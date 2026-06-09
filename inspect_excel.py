import openpyxl
from openpyxl.utils import get_column_letter

try:
    wb = openpyxl.load_workbook(r"c:\Users\heliu\OneDrive - Intel Corporation\Desktop\VSpractice\random-number-ui-app\docs\PCIE6_TEST_CASES.xlsx")
    print("Sheets:", wb.sheetnames)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n=== Sheet: {sheet_name} ===")
        print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")
        
        # Merged cells
        print("Merged cells:", list(ws.merged_cells.ranges)[:20])
        
        # Column widths
        for col in range(1, ws.max_column+1):
            letter = get_column_letter(col)
            dim = ws.column_dimensions.get(letter)
            if dim:
                print(f"  Col {letter}: width={dim.width}")
        
        # Row heights
        for row in range(1, min(ws.max_row+1, 10)):
            dim = ws.row_dimensions.get(row)
            if dim:
                print(f"  Row {row}: height={dim.height}")
        
        # Print first 15 rows with styles
        for row_idx in range(1, min(ws.max_row+1, 16)):
            for col_idx in range(1, ws.max_column+1):
                c = ws.cell(row=row_idx, column=col_idx)
                if c.value is not None:
                    fill_color = None
                    if c.fill and hasattr(c.fill, "fgColor") and c.fill.fgColor:
                        try:
                            fill_color = getattr(c.fill.fgColor, "rgb", str(getattr(c.fill.fgColor, "theme", "Unknown")))
                        except:
                            fill_color = "error"
                    font_info = None
                    if c.font:
                        font_info = f"bold={c.font.bold}, size={c.font.size}, color={c.font.color.rgb if c.font.color else None}"
                    print(f"  [{row_idx},{col_idx}] value={repr(c.value)[:80]}, fill={fill_color}, font={font_info}, align={c.alignment.horizontal if c.alignment else None}, wrap={c.alignment.wrap_text if c.alignment else None}")
except Exception as e:
    print(f"Error: {e}")
